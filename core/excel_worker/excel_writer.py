from openpyxl.styles import Font, Alignment
from core.excel_worker.workbook_manager import WorkbookManager


class ExcelWriter:
    @staticmethod
    def create_header(sheet, headers):
        header_font = Font(bold=True)
        for col_index, header in enumerate(headers, start=1):
            cell = sheet.cell(row=1, column=col_index)
            cell.value = header
            cell.font = header_font
            cell.alignment = Alignment(horizontal="center", vertical="center")

    @staticmethod
    def export_values(data_list, method_name):
        wb = WorkbookManager.get_workbook()
        if method_name in wb.sheetnames:
            print(f"Sheet '{method_name}' already exists")
            return
        sheet = wb.create_sheet(title=method_name)

        headers = ["X", "U", "V"]
        ExcelWriter.create_header(sheet, headers)

        for i, data in enumerate(data_list, start=2):
            d = data.as_dict()
            sheet.cell(row=i, column=1, value=d["x"])
            sheet.cell(row=i, column=2, value=d["u"])
            sheet.cell(row=i, column=3, value=d["v"])

    @staticmethod
    def export_functional(values, sheet_name):
        wb = WorkbookManager.get_workbook()
        if sheet_name in wb.sheetnames:
            print(f"Sheet '{sheet_name}' already exists")
            return
        sheet = wb.create_sheet(title=sheet_name)

        headers = ["Iteration", "J(v)"]
        ExcelWriter.create_header(sheet, headers)

        for i, val in enumerate(values, start=2):
            sheet.cell(row=i, column=1, value=i-2)
            sheet.cell(row=i, column=2, value=val)

    @staticmethod
    def export_gradient(gradient, sheet_name):
        wb = WorkbookManager.get_workbook()
        if sheet_name in wb.sheetnames:
            print(f"Sheet '{sheet_name}' already exists")
            return
        sheet = wb.create_sheet(title=sheet_name)

        headers = ["Iteration", "J'(v)"]
        ExcelWriter.create_header(sheet, headers)

        for i, val in enumerate(gradient, start=2):
            sheet.cell(row=i, column=1, value=i-2)
            sheet.cell(row=i, column=2, value=val)
